"""Export generated narrative to Word document using docxtpl."""

from datetime import date
from pathlib import Path
from io import BytesIO

from docxtpl import DocxTemplate

from schemas.clinical_data import ExtractedClinicalData
from schemas.narrative import AnthemReauthNarrative


def create_template_if_missing():
    """Create a basic template programmatically if the manual template doesn't exist."""
    template_path = Path(__file__).parent.parent / "templates" / "anthem_reauth.docx"
    if template_path.exists():
        return template_path

    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading("ABA Reauthorization Request", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Anthem Blue Cross California").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # Header fields
    doc.add_paragraph("Client ID: {{ client_id }}")
    doc.add_paragraph("Diagnosis: {{ diagnosis }}")
    doc.add_paragraph("Authorization Period: {{ auth_period_start }} — {{ auth_period_end }}")
    doc.add_paragraph("Supervising BCBA: {{ supervising_bcba }} ({{ bcba_credential }})")
    doc.add_paragraph("Date Prepared: {{ date_prepared }}")
    doc.add_paragraph("")

    # Sections
    sections = [
        ("1. Client Overview", "{{ client_overview }}"),
        ("2. Assessment Summary", "{{ assessment_summary }}"),
        ("3. Skill Acquisition Progress", "{{ skill_acquisition_progress }}"),
        ("4. Behavior Reduction Progress", "{{ behavior_reduction_progress }}"),
        ("5. Treatment Plan Update", "{{ treatment_plan_update }}"),
        ("6. Hours Utilization & Justification", "{{ hours_justification }}"),
        ("7. Caregiver Training", "{{ caregiver_training_summary }}"),
        ("8. Medical Necessity Statement", "{{ medical_necessity }}"),
        ("9. Discharge Plan", "{{ discharge_plan }}"),
    ]

    for heading, placeholder in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(placeholder)
        doc.add_paragraph("")

    # Footer
    doc.add_paragraph("")
    doc.add_paragraph("Prepared by: {{ supervising_bcba }}")
    doc.add_paragraph("Credential: {{ bcba_credential }}")
    doc.add_paragraph("Date: {{ date_prepared }}")
    doc.add_paragraph("")
    disclaimer = doc.add_paragraph(
        "DISCLOSURE: This document was drafted with AI assistance and reviewed/approved "
        "by the supervising BCBA listed above."
    )
    disclaimer.runs[0].font.size = Pt(9)

    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(template_path))
    return template_path


def export_to_docx(
    data: ExtractedClinicalData,
    narrative: AnthemReauthNarrative,
) -> bytes:
    """
    Render the narrative into a Word document.
    Returns the document as bytes for download.
    """
    template_path = create_template_if_missing()
    doc = DocxTemplate(str(template_path))

    context = {
        "client_id": data.client.client_id or "[CLIENT ID]",
        "diagnosis": ", ".join(
            f"{code} ({desc})"
            for code, desc in zip(
                data.client.diagnosis_codes,
                data.client.diagnosis_descriptions or data.client.diagnosis_codes,
            )
        ) or "[DIAGNOSIS]",
        "auth_period_start": str(data.client.auth_period_start or "[START DATE]"),
        "auth_period_end": str(data.client.auth_period_end or "[END DATE]"),
        "supervising_bcba": data.client.supervising_bcba or "[BCBA NAME]",
        "bcba_credential": data.client.bcba_credential_number or "[CREDENTIAL #]",
        "date_prepared": date.today().isoformat(),
        # Narrative sections
        "client_overview": narrative.client_overview.content,
        "assessment_summary": narrative.assessment_summary.content,
        "skill_acquisition_progress": narrative.skill_acquisition_progress.content,
        "behavior_reduction_progress": narrative.behavior_reduction_progress.content,
        "treatment_plan_update": narrative.treatment_plan_update.content,
        "hours_justification": narrative.hours_justification.content,
        "caregiver_training_summary": narrative.caregiver_training_summary.content,
        "medical_necessity": narrative.medical_necessity.content,
        "discharge_plan": narrative.discharge_plan.content,
    }

    doc.render(context)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
